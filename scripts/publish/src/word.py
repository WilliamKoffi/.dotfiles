from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


class Word:
    """Stateless namespace for rendering HTML to formatted DOCX."""

    COLOR_PRIMARY = RGBColor(2, 132, 199)
    COLOR_PRIMARY_DARK = RGBColor(3, 105, 161)
    COLOR_TEXT = RGBColor(15, 23, 42)
    COLOR_LINK = RGBColor(2, 132, 199)
    HEX_PRIMARY, HEX_BG_CODE, HEX_BG_ALT, HEX_BORDER = "0284C7", "F1F5F9", "F8FAFC", "E2E8F0"
    FONT_MAIN, FONT_HEADING, FONT_CODE = "Calibri", "Calibri", "Consolas"

    @staticmethod
    def render(html: str, target: str | Path) -> None:
        doc = DocxDocument()
        Word._setup(doc)
        soup = BeautifulSoup(html, "html.parser")
        for el in soup.children:
            if isinstance(el, NavigableString) and str(el).strip():
                doc.add_paragraph().add_run(str(el).strip())
            elif isinstance(el, Tag):
                t = el.name.lower()
                if t in ["h1", "h2", "h3", "h4", "h5", "h6"]: Word._heading(doc, el, int(t[1]))
                elif t == "p": Word._paragraph(doc, el)
                elif t == "pre":
                    code = el.find("code")
                    Word._box(doc, code.get_text() if code else el.get_text(), Word.HEX_BG_CODE)
                elif t == "blockquote": Word._box(doc, el, Word.HEX_BG_ALT, italic=True)
                elif t in ["ul", "ol"]: Word._list(doc, el, ordered=(t == "ol"))
                elif t == "table": Word._table(doc, el)
                elif t == "hr": Word._divider(doc)
        out = Path(target).resolve()
        if out.suffix.lower() in [".md", ".markdown"]:
            raise ValueError(f"Cannot render DOCX to a Markdown file: '{out.name}'")

        out.parent.mkdir(parents=True, exist_ok=True)

        tmp_out = out.with_name(f".{out.name}.tmp")
        try:
            doc.save(str(tmp_out))
            tmp_out.replace(out)
        except Exception:
            if tmp_out.exists():
                tmp_out.unlink()
            raise

    @staticmethod
    def _setup(doc: DocxDocument) -> None:
        for s in doc.sections:
            s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)
            s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = Word.FONT_MAIN, Pt(11)
        doc.styles["Normal"].font.color.rgb = Word.COLOR_TEXT

    @staticmethod
    def _paragraph(doc: DocxDocument, el: Tag) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(0), Pt(6), 1.15
        Word._inline(el, p)

    @staticmethod
    def _inline(el: Tag, p, bold=False, italic=False, code=False) -> None:
        for c in el.children:
            if isinstance(c, NavigableString):
                if not str(c): continue
                run = p.add_run(str(c))
                run.bold, run.italic, run.font.name = bold, italic, (Word.FONT_CODE if code else Word.FONT_MAIN)
                if code: run.font.size, run.font.color.rgb = Pt(9.5), Word.COLOR_PRIMARY_DARK
            elif isinstance(c, Tag):
                t = c.name.lower()
                if t in ["strong", "b"]: Word._inline(c, p, True, italic, code)
                elif t in ["em", "i"]: Word._inline(c, p, bold, True, code)
                elif t == "code": Word._inline(c, p, bold, italic, True)
                elif t == "a":
                    r = p.add_run(c.get_text())
                    r.font.name, r.font.color.rgb, r.underline = Word.FONT_MAIN, Word.COLOR_LINK, True
                elif t == "kbd":
                    r = p.add_run(f"[{c.get_text()}]")
                    r.font.name, r.bold = Word.FONT_CODE, True
                elif t == "br": p.add_run("\n")
                else: Word._inline(c, p, bold, italic, code)

    @staticmethod
    def _heading(doc: DocxDocument, tag: Tag, level: int) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = {1: Pt(18), 2: Pt(14), 3: Pt(10)}.get(level, Pt(6))
        p.paragraph_format.space_after = {1: Pt(8), 2: Pt(6), 3: Pt(4)}.get(level, Pt(3))
        Word._inline(tag, p, bold=True)
        for r in p.runs:
            r.font.name, r.font.size = Word.FONT_HEADING, {1: Pt(22), 2: Pt(15), 3: Pt(12.5)}.get(level, Pt(11))
            r.font.color.rgb = {1: Word.COLOR_PRIMARY, 2: Word.COLOR_PRIMARY_DARK}.get(level, Word.COLOR_TEXT)

    @staticmethod
    def _box(doc: DocxDocument, content: str | Tag, bg: str, italic: bool = False) -> None:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>'))
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="0" w:color="{Word.HEX_PRIMARY}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
        if isinstance(content, Tag): Word._inline(content, p, italic=italic)
        else:
            p.paragraph_format.line_spacing = 1.05
            lines = content.strip("\n").split("\n")
            for i, line in enumerate(lines):
                r = p.add_run(line)
                r.font.name, r.font.size, r.font.color.rgb = Word.FONT_CODE, Pt(9.0), Word.COLOR_TEXT
                if i < len(lines) - 1: p.add_run("\n")
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    @staticmethod
    def _list(doc: DocxDocument, tag: Tag, ordered=False, level=0) -> None:
        for idx, li in enumerate(tag.find_all("li", recursive=False), start=1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.line_spacing = Inches(0.25 * (level + 1)), 1.15
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(1), Pt(2)
            p.add_run(f"{idx}. " if ordered else "▪ ").bold = True
            p.runs[0].font.color.rgb = Word.COLOR_PRIMARY
            for c in li.children:
                if isinstance(c, Tag) and c.name in ["ul", "ol"]: Word._list(doc, c, ordered=(c.name == "ol"), level=level + 1)
                elif isinstance(c, Tag): Word._inline(c, p)
                elif isinstance(c, NavigableString) and str(c).strip(): p.add_run(str(c)).font.name = Word.FONT_MAIN

    @staticmethod
    def _table(doc: DocxDocument, tag: Tag) -> None:
        rows = tag.find_all("tr")
        if not rows: return
        cols = max(len(r.find_all(["th", "td"])) for r in rows)
        if cols == 0: return
        table = doc.add_table(rows=len(rows), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table._tbl.tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{Word.HEX_BORDER}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{Word.HEX_PRIMARY}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{Word.HEX_BORDER}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>'))
        for r_idx, r in enumerate(rows):
            cells = r.find_all(["th", "td"])
            is_head = (r_idx == 0 and r.find("th") is not None) or all(c.name == "th" for c in cells)
            for c_idx, c_tag in enumerate(cells[:cols]):
                cell, p = table.cell(r_idx, c_idx), table.cell(r_idx, c_idx).paragraphs[0]
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                bg = Word.HEX_PRIMARY if is_head else (Word.HEX_BG_ALT if r_idx % 2 == 1 else None)
                if bg: cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>'))
                Word._inline(c_tag, p, bold=is_head)
                for run in p.runs:
                    if is_head: run.font.color.rgb, run.font.size = RGBColor(255, 255, 255), Pt(10)
                    else: run.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    @staticmethod
    def _divider(doc: DocxDocument) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("―" * 45).font.color.rgb = RGBColor(203, 213, 225)
